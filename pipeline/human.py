#!/usr/bin/env python3
"""
Packets for the researchers' steps. The pipeline never does these steps itself; it prepares what
the researchers need and reads back what they return.

  blind-pass        human/01_blind_pass/: the seeded, stratified sample to open-code BEFORE anyone sees
                    model output (sheet or CSV, plus instructions)
  review-dashboard  human/02_author_review/codebook_review.html: one self-contained page to review every
                    candidate code against every unit it cites, with the council's warning signs, the
                    challenge logs (Full) and the quarantined candidates; decisions are downloaded as JSON
  review-packet     the same review as a workbook (xlsx when openpyxl is installed, else CSV files)
  review-to-edits   compare two or more downloaded review files; write the agreements as a draft edit log
                    (codebooks/review_edits.draft.json) and the disagreements to negotiate
  heldout-packet    human/03_heldout_coding/: the held-out evaluation units x the approved codes, no model
                    assignment anywhere; the researchers code it independently and resolve differences
  spot-check        human/05_spot_check/: issue-coded units, uncovered content, a random sample of consensus
                    assignments with rationales, and (Full) the decision model's disagreements
  low-alpha         human/06_low_alpha/: every code with Krippendorff's alpha below the floor, with its
                    definition and the units the coders disagreed on; the researchers return decisions.csv
                    (refine and recode, or drop), which consensus, the report and the dashboard read

Standard library (+ openpyxl if installed, for .xlsx).
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import sys
from collections import OrderedDict, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import council as K  # noqa: E402

TEMPLATE = Path(__file__).resolve().parent / "templates" / "review_dashboard.html"


# --------------------------------------------------------------------------- workbook writer
def write_book(base: Path, sheets: "OrderedDict[str, tuple]"):
    """sheets: name -> (header, rows, widths). Writes base.xlsx if openpyxl is available; always
    writes one CSV per sheet in base_csv/."""
    base.parent.mkdir(parents=True, exist_ok=True)
    cdir = base.parent / (base.name + "_csv")
    cdir.mkdir(exist_ok=True)
    for name, (header, rows, _) in sheets.items():
        with open(cdir / ("%s.csv" % name.replace(" ", "_")), "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(header)
            w.writerows(rows)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("   (openpyxl not installed: CSV files only, in %s)" % cdir)
        return cdir
    wb = Workbook()
    wb.remove(wb.active)
    for name, (header, rows, widths) in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="305496")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append([str(x)[:32000] if x is not None else "" for x in r])
        for i, wdt in enumerate(widths or []):
            ws.column_dimensions[get_column_letter(i + 1)].width = wdt
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
    out = base.with_suffix(".xlsx")
    wb.save(out)
    return out


def unit_text(r) -> str:
    return r["text"] + (("\n\nCONTEXT: " + r["context"]) if (r.get("context") or "").strip() else "")


# --------------------------------------------------------------------------- blind pass
BLIND_README = """# Blind pass (open coding before any model output)

**Who:** at least two researchers, each working alone.
**Return:** your filled copy of the blind-pass sheet, renamed with your initials.

## Why

The council proposes codes from the whole corpus, and author review then judges those codes
against the units each one cites. Review can only judge what the models chose to propose, and
model output anchors reviewers. This blind pass is the only check on what the models did not
select: you code a sample yourself before seeing anything the models produced. After the review,
you list which of your codes the approved codebook lacks.

## The sample

{n} {units}, drawn with seed {seed}, stratified by {strata} ({note}).

## What to do

1. **Do not look at any model output first**: no discovery file, codebook, challenge log or
   dashboard. Do not discuss codes with the other researchers until everyone has returned a file.
2. Read each unit in `Code here`. {scope}
3. Code against the research questions:
{rqs}
4. Write short code labels in `your codes`, separated by `;`, in your own words, creating codes as
   you go. A unit can have none, one or several. Put the question ids each code answers in
   `research questions`.
5. List your codes in `My codes`, each with a one-sentence meaning and one example unit id.
6. Fill in `About you`, including whether you had seen any model output or prior coding of these units.

There are no right answers. The point is your independent reading.
"""


def cmd_blind_pass(args):
    corpus = K.load_corpus()
    sp = K.splits()
    units = sp.get("blind_pass") or []
    s = K.cfg.study
    has_link = any(r.get("link") for r in corpus.values())
    rows = [[u] + ([corpus[u].get("link", "")] if has_link else []) + [unit_text(corpus[u]), "", ""] for u in units]
    out = write_book(K.D("human", "01_blind_pass") / "blind_pass_TEMPLATE", OrderedDict([
        ("Code here", (["uid"] + (["link"] if has_link else []) + ["text", "your codes (; separated)", "research questions"], rows, [8] + ([16] if has_link else []) + [90, 40, 14])),
        ("My codes", (["code label", "one-sentence meaning", "research question(s)", "example uid"], [], [35, 60, 14, 10])),
        ("About you", (["question", "answer"], [["Your name or initials", ""], ["Date", ""], ["Minutes spent", ""],
                                                ["Had you seen any model output or prior coding of these units before finishing? (describe)", ""]], [80, 40])),
    ]))
    strata = ", ".join(c for c in K.cfg.d["splits"]["strata_columns"]) or "none"
    (K.W / "human" / "01_blind_pass" / "README.md").write_text(BLIND_README.format(
        n=len(units), units=s["unit_name_plural"], seed=sp.get("seed"), strata=strata, note=sp.get("blind_pass_note", ""),
        scope=(s.get("scope_rule") or "").strip(), rqs="\n".join("   - **%s** %s" % (k, q["text"]) for k, q in K.cfg.rqs.items())), encoding="utf-8")
    print("blind-pass: %d units -> %s (+ README.md). Hand it to at least two researchers BEFORE they see any model output." % (len(units), out))
    return 0


# --------------------------------------------------------------------------- review signals
def review_signals(book: dict, cond: str):
    """Warning signs from the council's own artifacts only (never from any external coding)."""
    W = K.W
    ground, rejected = {}, {}
    for L in K.cfg.analysts:
        for x in K.rj(W / "discovery" / "jev" / ("grounding_%s.json" % L)).get("all", []):
            ground.setdefault("%s:%s" % (L, x["code_id"]), {})[x["uid"]] = x["p"]
        chs = {c.get("id"): c for c in K.rj(W / "discovery" / "challenges" / ("adversary_%s.json" % L)).get("challenges", [])}
        for a in K.rj(W / "discovery" / "answers" / ("answer_%s.json" % L)).get("answers", []):
            ch = chs.get(a.get("challenge_id"), {})
            if a.get("answer") == "REJECT" and ch.get("code_id"):
                rejected.setdefault("%s:%s" % (L, ch["code_id"]), []).append(
                    {"id": a.get("challenge_id"), "type": ch.get("type"), "argument": ch.get("argument", ""), "reason": a.get("reason", "")})
    prov = book.get("provenance", {})
    answers = {a.get("challenge_id"): a for a in K.rj(W / "reconciliation" / cond / "reconciler_answers.json").get("answers", [])}
    recon = {}
    for c in K.rj(W / "reconciliation" / cond / "adversary_challenges.json").get("challenges", []):
        a = answers.get(c.get("id"), {})
        item = {"id": c.get("id"), "type": c.get("type"), "evidence": c.get("evidence", ""), "argument": c.get("argument", ""),
                "answer": a.get("answer", ""), "reason": a.get("reason", ""), "revision": a.get("revision", "")}
        for m in c.get("members", []):
            for cid in prov.get(m, {}).get("codes", []):
                recon.setdefault(cid, []).append(item)
    return ground, rejected, recon


def code_record(c, ground, rejected, recon):
    keys = [s["key"] for s in c.get("source_codes", [])]
    gp = {}
    for k in keys:
        for u, p in ground.get(k, {}).items():
            if p is not None:
                gp[u] = min(p, gp.get(u, 1.0))
    lo = float(K.cfg.t["grounding_lead_below"])
    weak = [u for u in c["units"] if gp.get(u, 1.0) < lo]
    unit = K.cfg.study["unit_name_plural"]
    flags = []
    if c["n_analysts"] <= 1:
        flags.append("Proposed by one analyst")
    if gp and len(weak) / max(1, len(c["units"])) >= 0.5:
        flags.append("Weak grounding (%d of %d %s below %s)" % (len(weak), len(c["units"]), unit, lo))
    rej = [x for k in keys for x in rejected.get(k, [])]
    if rej:
        flags.append("Rejected discovery challenge")
    if [x for x in recon.get(c["id"], []) if x["answer"] != "ACCEPT"]:
        flags.append("Reconciliation challenge not fully accepted")
    if c["n_units"] <= int(K.cfg.t["support_floor"]) + 1:
        flags.append("Thin support (%d %s)" % (c["n_units"], unit))
    return OrderedDict([
        ("id", c["id"]), ("rq", c.get("rq", "")), ("label", c["label"]), ("definition", c["definition"]),
        ("include", c.get("include", "")), ("exclude", c.get("exclude", "")), ("n_units", c["n_units"]),
        ("n_analysts", c["n_analysts"]), ("layers", c.get("layers", {})),
        ("sources", [{"key": s["key"], "label": s["label"], "layer": s.get("layer", "")} for s in c.get("source_codes", [])]),
        ("quotes", [{"uid": q.get("uid"), "quote": q.get("quote")} for q in c.get("anchor_quotes", [])][:6]),
        ("units", c["units"]), ("weak", weak), ("flags", flags),
        ("priority", "full" if len(flags) >= 2 else "check" if flags else "quick"),
        ("rejected", rej), ("recon", recon.get(c["id"], []))])


def candidate(args) -> tuple:
    cond = getattr(args, "cond", None) or K.cfg.mode
    p = K.W / "codebooks" / ("candidate_%s.json" % cond)
    if not p.exists():
        raise SystemExit("no candidate codebook %s; run `run.py reconcile` first" % p.relative_to(K.W))
    return K.read_json(p), cond


def cmd_review_dashboard(args):
    corpus = K.load_corpus()
    sp = K.splits()
    dev = set(sp.get("development") or [])
    book, cond = candidate(args)
    ground, rejected, recon = review_signals(book, cond)
    codes = [code_record(c, ground, rejected, recon) for c in book["codes"]]
    used = {u for c in codes for u in c["units"]}
    quarantined = [{"key": q["key"], "label": q.get("label", ""), "definition": q.get("definition", ""),
                    "units": q.get("units", []), "assessment": q.get("reconciler_assessment")} for q in book.get("quarantined", [])]
    for q in quarantined:
        used.update(q["units"])
    used.update(dev)
    data = {"rqs": OrderedDict((k, q["text"]) for k, q in K.cfg.rqs.items()), "codes": codes,
            "relations": [{"type": r["type"], "from": r.get("from_code"), "to": r.get("to_code"), "note": r.get("note", "")}
                          for r in book.get("relations", []) if r.get("from_code") and r.get("to_code")],
            "quarantined": quarantined,
            "threads": {u: {"link": corpus[u].get("link", ""), "text": unit_text(corpus[u]), "dev": u in dev} for u in corpus if u in used},
            "dev": [u for u in corpus if u in dev], "n_units": len(corpus), "unit_plural": K.cfg.study["unit_name_plural"],
            "mode": "%s, candidate codebook %s" % (K.cfg.mode.capitalize(), cond), "round": args.round,
            "proposals": K.rj(Path(args.proposals)) if args.proposals else {},
            "store": "council-review:%s" % hashlib.sha256((str(K.W) + json.dumps([c["id"] for c in codes]) +
                                                         K.read_json(K.W / "manifest.json").get("corpus_txt_sha256", "")).encode()).hexdigest()[:16]}
    if not codes:
        raise SystemExit("the candidate codebook has no codes")
    html = TEMPLATE.read_text(encoding="utf-8").replace("/*__DATA__*/null", json.dumps(data, ensure_ascii=False).replace("</", "<\\/"))
    out = K.D("human", "02_author_review", "codebook_review.html")
    out.write_text(html, encoding="utf-8")
    (out.parent / "README.md").write_text(REVIEW_README.format(n=len(codes), unit=K.cfg.study["unit_name_plural"],
                                                                floor=K.cfg.t["support_floor"]), encoding="utf-8")
    print("review-dashboard: %s (%d codes, %d units embedded, %d KB); priorities full/check/quick = %d/%d/%d. "
          "STOP here: researchers review (they must have finished the blind pass first)."
          % (out, len(codes), len(data["threads"]), len(html) // 1024, sum(c["priority"] == "full" for c in codes),
             sum(c["priority"] == "check" for c in codes), sum(c["priority"] == "quick" for c in codes)))
    return 0


REVIEW_README = """# Author review and codebook approval

**Who:** at least two researchers. Each reviews alone first, then you meet and agree the edits.
**Before you start:** finish the blind pass. This page shows model output.
**Open:** `codebook_review.html` in a browser (it is self-contained; nothing is sent anywhere).

## What you approve

{n} candidate codes. Once approved, three coders apply the codebook to every unit and nothing in
it changes during coding.

## For each code

- **Fidelity:** does the label say what the {unit} say?
- **Grounding:** do the cited {unit} support it? Mark any that do not.
- **Distinctness:** could a coder tell it from its neighbours using the definitions alone?

Then decide: keep, relabel or redefine, merge, move to another question, split, or drop. Add one or
two positive and one or two negative examples, **only from the development units** (marked
example-eligible): examples drawn from the units later used to evaluate coding inflate agreement.
In Full, read the challenge logs: a rejected challenge is a disagreement between two models for
you to settle. Admit or reject each adversary-proposed (quarantined) code on its own tab.

Priority comes only from the council's warning signs: proposed by one analyst; weak grounding
(the decision model scored most cited units below the threshold, Full only); a rejected discovery
challenge; a reconciliation challenge not fully accepted; thin support (at most {floor} plus one).

## When you finish

1. `Finish review` tab: enter your name, list the blind-pass codes the codebook lacks, and
   **Download my decisions** (a JSON file).
2. Run `python3 pipeline/human.py --config CONFIG review-to-edits review_A.json review_B.json`.
   It writes the edits you agree on to `codebooks/review_edits.draft.json` and lists the
   disagreements in `human/02_author_review/review_comparison.md`.
3. Meet, resolve every disagreement by negotiated agreement, and save the final, agreed edit log as
   `codebooks/review_edits.json` (operations and fields: AGENTS.md, Stage 3).
   Every edit has a reason. An added code cites at least {floor} units.
4. `python3 pipeline/council.py --config CONFIG approve` closes and versions the codebook.

If review finds a gap that no candidate and its evidence can fill, discovery reruns with your
notes (AGENTS.md, Stage 3).
"""


def cmd_review_packet(args):
    corpus = K.load_corpus()
    dev = set(K.splits().get("development") or [])
    book, cond = candidate(args)
    ground, rejected, recon = review_signals(book, cond)
    recs = [code_record(c, ground, rejected, recon) for c in book["codes"]]
    rows = [[r["priority"], "; ".join(r["flags"]), r["rq"], r["id"], r["label"], r["definition"], r["include"], r["exclude"], r["n_units"],
             r["n_analysts"], "; ".join("%s (%s): %s" % (s["key"], s["layer"], s["label"]) for s in r["sources"]),
             "\n".join('%s: "%s"' % (q["uid"], q["quote"]) for q in r["quotes"]), "", "", "", "", "", "", ""] for r in recs]
    chs = K.rj(K.W / "reconciliation" / cond / "adversary_challenges.json").get("challenges", [])
    ans = {a.get("challenge_id"): a for a in K.rj(K.W / "reconciliation" / cond / "reconciler_answers.json").get("answers", [])}
    out = write_book(K.D("human", "02_author_review") / "author_review_TEMPLATE", OrderedDict([
        ("Codes", (["priority", "why flagged", "rq", "id", "label", "definition", "include", "exclude", "n units", "n analysts", "source codes",
                    "anchor quotes", "FIDELITY ok? (y/n + note)", "GROUNDING ok? (y/n + note)", "DISTINCTNESS ok? (y/n + note)",
                    "positive example uids (development only)", "negative example uids (development only)", "decision (keep/relabel/merge/move/split/drop)", "reason"],
                   rows, [10, 30, 8, 6, 30, 40, 30, 30, 7, 7, 40, 50, 18, 18, 18, 18, 18, 22, 30])),
        ("Cited units", (["code", "uid", "split", "text", "supports the code? (y/n + note)"],
                         [[c["id"], u, "development" if u in dev else "evaluation", unit_text(corpus[u]), ""] for c in book["codes"] for u in c["units"]],
                         [6, 8, 11, 100, 30])),
        ("Recon challenges", (["id", "type", "group", "members", "evidence", "argument", "reconciler answer", "reconciler reason", "YOUR DECISION"],
                              [[c.get("id"), c.get("type"), c.get("group_label"), ", ".join(c.get("members", [])), c.get("evidence", ""), c.get("argument", ""),
                                ans.get(c.get("id"), {}).get("answer", ""), ans.get(c.get("id"), {}).get("reason", ""), ""] for c in chs],
                              [8, 18, 30, 25, 40, 50, 12, 40, 35])),
        ("Quarantined", (["key", "label", "definition", "n units", "units", "reconciler assessment", "ADMIT? (y/n)", "reason"],
                         [[q.get("key"), q.get("label"), q.get("definition"), q.get("n_units"), ", ".join(q.get("units", [])),
                           json.dumps(q.get("reconciler_assessment")), "", ""] for q in book.get("quarantined", [])], [10, 30, 45, 7, 30, 35, 10, 30])),
        ("Example units", (["uid", "text"], [[u, unit_text(corpus[u])[:3000]] for u in corpus if u in dev], [8, 110])),
        ("Edit log", (["reviewer(s)", "op", "code id(s)", "units (move/add/split)", "new label / definition / include / exclude", "reason"], [], [14, 20, 14, 20, 60, 50])),
    ]))
    print("review-packet: %s (%d codes)" % (out, len(recs)))
    return 0


# --------------------------------------------------------------------------- review files -> draft edit log
def cmd_review_to_edits(args):
    book, cond = candidate(args)
    files = [Path(f) for f in args.files] or sorted((K.W / "human" / "02_author_review" / "returned").glob("*.json"))
    if not files:
        raise SystemExit("no review files: pass them, or put each reviewer's downloaded review_<name>.json in "
                         "human/02_author_review/returned/")
    revs = OrderedDict()
    for f in files:
        d = json.loads(f.read_text(encoding="utf-8"))
        revs[d.get("reviewer") or f.stem] = d
    if len(revs) < 2:
        print("   ! only %d review file(s); the method requires at least two researchers" % len(revs))
    by_id = {c["id"]: c for c in book["codes"]}
    agree, negotiate, lines = [], [], ["# Author review: comparison of %d reviewers (%s)" % (len(revs), ", ".join(revs)), ""]

    def decision(r, cid):
        x = (r.get("reviews") or {}).get(cid) or {}
        return OrderedDict([("decision", x.get("decision") or "keep"), ("mergeInto", x.get("mergeInto")), ("moveTo", x.get("moveTo")),
                            ("newLabel", (x.get("newLabel") or "").strip()), ("newDefinition", (x.get("newDefinition") or "").strip()),
                            ("pos", x.get("pos") or []), ("neg", x.get("neg") or []), ("unsupported", x.get("unsupported") or []),
                            ("note", (x.get("note") or "").strip()),
                            ("criteria", {k: x.get(k) for k in ("fidelity", "grounding", "distinctness")})])

    lines += ["| Code | Label | " + " | ".join(revs) + " | Agree? |", "|---|---|" + "---|" * (len(revs) + 1)]
    for cid, c in by_id.items():
        ds = OrderedDict((n, decision(r, cid)) for n, r in revs.items())
        kinds = {(d["decision"], d["mergeInto"], d["moveTo"]) for d in ds.values()}
        same = len(kinds) == 1
        lines.append("| %s | %s | %s | %s |" % (cid, c["label"], " | ".join("%s%s%s" % (d["decision"], (" into " + d["mergeInto"]) if d["mergeInto"] else "",
                                                                                   (" to " + d["moveTo"]) if d["moveTo"] else "") for d in ds.values()),
                                                 "yes" if same else "**negotiate**"))
        first = next(iter(ds.values()))
        reasons = "; ".join("%s: %s" % (n, d["note"]) for n, d in ds.items() if d["note"])
        pos = sorted(set.intersection(*[set(d["pos"]) for d in ds.values()])) if ds else []
        neg = sorted(set.intersection(*[set(d["neg"]) for d in ds.values()])) if ds else []
        if pos or neg:
            agree.append(OrderedDict([("op", "examples"), ("id", cid), ("positive", pos), ("negative", neg)]))
        unsup = sorted(set.union(*[set(d["unsupported"]) for d in ds.values()])) if ds else []
        if unsup:
            negotiate.append(OrderedDict([("id", cid), ("issue", "units marked as not supporting the code"), ("units", unsup)]))
        if not same:
            negotiate.append(OrderedDict([("id", cid), ("issue", "different decisions"), ("decisions", ds)]))
            continue
        k = first["decision"]
        if k == "drop":
            agree.append(OrderedDict([("op", "drop"), ("id", cid), ("reason", reasons or "both reviewers: drop")]))
        elif k == "merge" and first["mergeInto"]:
            agree.append(OrderedDict([("op", "merge"), ("into", first["mergeInto"]), ("from", [cid]), ("reason", reasons or "both reviewers: merge")]))
        elif k == "move" and first["moveTo"]:
            agree.append(OrderedDict([("op", "relabel"), ("id", cid), ("rq", first["moveTo"]), ("reason", reasons or "both reviewers: move to %s" % first["moveTo"])]))
        elif k in ("relabel", "split"):
            negotiate.append(OrderedDict([("id", cid), ("issue", "agreed to %s; agree the wording%s" % (k, " and the units of each part" if k == "split" else "")),
                                          ("proposals", OrderedDict((n, {"label": d["newLabel"], "definition": d["newDefinition"]}) for n, d in ds.items()))]))
    lines += ["", "## To negotiate (%d)" % len(negotiate), ""]
    for n in negotiate:
        lines.append("- **%s** %s: %s" % (n["id"], by_id[n["id"]]["label"], n["issue"]))
    q_dec = defaultdict(dict)
    for n, r in revs.items():
        for key, x in (r.get("reviews") or {}).items():
            if key.startswith("Q_") and x.get("decision"):
                q_dec[key[2:]][n] = (x["decision"], x.get("note", ""))
    for qk, ds in q_dec.items():
        vals = {d for d, _ in ds.values()}
        if len(vals) == 1 and len(ds) == len(revs):
            if "reject" in vals:
                agree.append(OrderedDict([("op", "reject_quarantined"), ("key", qk), ("reason", "; ".join(r.rstrip(". ") for _, r in ds.values() if r) or "rejected by all reviewers")]))
            else:
                agree.append(OrderedDict([("op", "accept_quarantined"), ("key", qk), ("id", "CHOOSE_A_NEW_ID"), ("reason", "; ".join(r.rstrip(". ") for _, r in ds.values() if r) or "admitted by all reviewers")]))
        else:
            negotiate.append(OrderedDict([("id", qk), ("issue", "quarantined candidate: reviewers differ or did not all decide"), ("decisions", ds)]))
    missing = sorted({m for r in revs.values() for m in r.get("blind_pass_missing") or []})
    draft = OrderedDict([("reviewed_by", ", ".join(revs)), ("status_note", "DRAFT from review-to-edits; resolve to_negotiate, then save as review_edits.json"),
                         ("blind_pass_codes_missing", missing), ("edits", agree), ("to_negotiate", negotiate)])
    K.write_json(K.D("codebooks", "review_edits.draft.json"), draft)
    (K.D("human", "02_author_review", "review_comparison.md")).write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("review-to-edits: %d agreed edits, %d items to negotiate -> codebooks/review_edits.draft.json, human/02_author_review/review_comparison.md"
          % (len(agree), len(negotiate)))
    return 0


# --------------------------------------------------------------------------- held-out coding
HELDOUT_README = """# Held-out coding (correctness check)

**Who:** at least two researchers, independently. **Do not look at any model assignment** for these units.

Krippendorff's alpha measures consistency among the model coders, not correctness. This sample
is the correctness check: {n} {units} drawn at random from the evaluation units (seed {seed}),
coded against the approved codebook (version {version}).

1. Each researcher codes every row of `heldout_coding_TEMPLATE` alone: 1 in the column of every code
   that applies, an issue code where its definition is met, and a note under UNCOVERED for relevant
   content no code fits. Use the definitions, include and exclude rules on the `Codebook` sheet.
2. Compare, resolve every difference by negotiated agreement, and save ONE resolved file as CSV with
   the same columns (`uid` and one 0/1 column per code id).
3. `python3 pipeline/council.py --config CONFIG heldout-score --labels resolved.csv` scores council
   consensus against your resolved labels (Cohen's kappa per code and pooled).
"""


def cmd_heldout_packet(args):
    corpus = K.load_corpus()
    sp = K.splits()
    book = K.approved_book()
    units = sp.get("held_out") or []
    codes = book["codes"]
    issues = list(K.cfg.issue_codes)
    header = ["uid", "text"] + ["%s %s" % (c["id"], c["label"][:40]) for c in codes] + issues + ["UNCOVERED note"]
    rows = [[u, unit_text(corpus[u])] + [""] * (len(codes) + len(issues) + 1) for u in units]
    cb = [[c["id"], c.get("rq", ""), c["label"], c["definition"], c.get("include", ""), c.get("exclude", ""),
           ", ".join((c.get("examples") or {}).get("positive", [])), ", ".join((c.get("examples") or {}).get("negative", []))] for c in codes]
    cb += [[k, "", "issue code", v, "", "", "", ""] for k, v in K.cfg.issue_codes.items()]
    out = write_book(K.D("human", "03_heldout_coding") / "heldout_coding_TEMPLATE", OrderedDict([
        ("Code here", (header, rows, [8, 90] + [9] * (len(codes) + len(issues)) + [30])),
        ("Codebook", (["id", "rq", "label", "definition", "include", "exclude", "positive examples", "negative examples"], cb, [8, 8, 35, 50, 40, 40, 16, 16])),
        ("About you", (["question", "answer"], [["Your name or initials", ""], ["Date", ""], ["Minutes spent", ""],
                                                ["Did you see any model assignment for these units before finishing? (should be no)", ""]], [80, 40])),
    ]))
    (K.W / "human" / "03_heldout_coding" / "README.md").write_text(HELDOUT_README.format(
        n=len(units), units=K.cfg.study["unit_name_plural"], seed=sp.get("seed"), version=book.get("version")), encoding="utf-8")
    print("heldout-packet: %d units x %d codes -> %s. Researchers code it without seeing model assignments." % (len(units), len(codes), out))
    return 0


# --------------------------------------------------------------------------- spot-check
def cmd_spot_check(args):
    corpus = K.load_corpus()
    book = K.approved_book()
    cond = args.cond or K.primary_condition()
    res = K.W / "results" / cond
    if not (res / "summary.json").exists():
        raise SystemExit("run `council.py consensus` first")
    summ = K.read_json(res / "summary.json")
    long = defaultdict(dict)
    for r in K.read_csv(res / "coder_long.csv"):
        long[r["uid"]][r["coder"]] = "%s :: %s" % (r["codes"], r["why"])
    labels = {c["id"]: c["label"] for c in book["codes"]}
    rat = lambda u: " || ".join("coder %s: %s" % (k, v) for k, v in sorted(long[u].items()))
    rows = []
    for u in summ["issue_units"]:
        rows.append(["issue (2+ coders)", u, "ISSUE", "", unit_text(corpus[u]), rat(u), "", ""])
    for u, notes in summ["uncovered_units"].items():
        rows.append(["uncovered", u, "UNCOVERED", " / ".join(dict.fromkeys(n["note"] for n in notes)), unit_text(corpus[u]), rat(u), "", ""])
    cons = [r for r in K.read_csv(res / "consensus.csv") if r["code"] in labels]
    scr = K.W / "coding" / "jev" / "screen.json"
    if K.cfg.full and scr.exists():
        sc = K.read_json(scr)["scores"]
        lo, hi = float(K.cfg.t["coding_lead_assigned_below"]), float(K.cfg.t["coding_lead_unassigned_at_or_above"])
        cset = {(r["uid"], r["code"]) for r in cons}
        for r in cons:
            p = sc.get(r["uid"], {}).get(r["code"])
            if p is not None and p < lo:
                rows.append(["uncertain: assigned at consensus, decision model p=%.2f (held from prevalence until kept)" % p, r["uid"], r["code"],
                             labels[r["code"]], unit_text(corpus[r["uid"]]), rat(r["uid"]), "", ""])
        for u, row in sc.items():
            for c, p in row.items():
                if p is not None and p >= hi and (u, c) not in cset and c in labels:
                    rows.append(["possible miss: not at consensus, decision model p=%.2f" % p, u, c, labels[c], unit_text(corpus[u]), rat(u), "", ""])
    rng = random.Random(int(K.cfg.d["spot_check"]["seed"]))
    sample = rng.sample(cons, min(int(K.cfg.d["spot_check"]["random_sample"]), len(cons)))
    for r in sample:
        rows.append(["random consensus assignment", r["uid"], r["code"], labels[r["code"]], unit_text(corpus[r["uid"]]), rat(r["uid"]), "", ""])
    d = K.D("human", "05_spot_check")
    with open(d / "spot_check.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["kind", "uid", "code", "label", "text", "coders' codes :: rationale", "keep", "note"])
        w.writerows(rows)
    (d / "README.md").write_text(
        "# Spot-check of rationales\n\nAt least two researchers check each row's rationales against the codebook and fill `keep` "
        "(y or n) and `note`. Report the sample size. Save the agreed file as `decisions.csv` in this folder (same columns); "
        "`council.py consensus` reads `uid`, `code` and `keep` from it. Uncertain assignments (Full) stay out of prevalence until "
        "a row keeps them. Issue-coded units and uncovered content are reported as findings about the data and the codebook.\n",
        encoding="utf-8")
    print("spot-check (%s): %d rows (%d random consensus assignments) -> %s" % (cond, len(rows), len(sample), d / "spot_check.csv"))
    return 0


# --------------------------------------------------------------------------- low-alpha decision
LOW_ALPHA_README = """# Codes below the reliability floor: refine and recode, or drop

**Who:** the researchers (at least two), together.

The three coders could not apply {n} code(s) consistently: Krippendorff's alpha on the
{rel} units is below {floor} in condition {cond}. A code below the floor cannot support a finding.
For each one, decide:

- **drop**: the code is excluded from the results and listed, with its definition and agreement
  counts, in the supplement (`results/{cond}/reliability_report.md`).
- **refine**: you rewrite the definition (and include / exclude) so coders can apply it; the
  codebook gets a new version and the WHOLE corpus is recoded with it (a code is never recoded on
  part of the corpus). Codes you drop are removed from the new version.

`low_alpha_codes.md` shows each code's definition, agreement counts and every unit the coders
disagreed on, with each coder's codes and rationale.

## When you have decided

1. Copy `decisions_TEMPLATE.csv` to `decisions.csv` in this folder and fill one row per code:
   `decision` is `drop` or `refine`; for `refine`, fill `new_definition` (and `new_label`,
   `new_include`, `new_exclude` where they change); give a `reason` and who decided.
2. `python3 pipeline/council.py --config CONFIG consensus` records the decisions in the reports.
3. If any code is refined: `python3 pipeline/council.py --config CONFIG refine` archives this coding
   round under `archive/`, writes the new codebook version, and then
   `python3 pipeline/run.py --config CONFIG code` recodes the corpus. Recode the held-out sample for
   the refined codes as well.
"""


def cmd_low_alpha(args):
    corpus = K.load_corpus()
    book = K.approved_book()
    cond = args.cond or K.primary_condition()
    res = K.W / "results" / cond
    if not (res / "summary.json").exists():
        raise SystemExit("run `council.py consensus` first")
    summ = K.read_json(res / "summary.json")
    low = summ["codes_dropped_alpha_below_floor"]
    undefined = summ["codes_alpha_undefined"]
    d = K.D("human", "06_low_alpha")
    books = {c["id"]: c for c in book["codes"]}
    rel = {r["code"]: r for r in K.read_csv(res / "reliability.csv")}
    long = defaultdict(dict)
    for r in K.read_csv(res / "coder_long.csv"):
        long[r["uid"]][r["coder"]] = (r["codes"].split("|") if r["codes"] else [], r["why"])
    rel_units = set(K.splits().get("evaluation") or corpus) if K.cfg.d["reliability"]["units"] != "all" else set(corpus)
    L = ["# Codes below alpha %s (condition %s, codebook version %s)" % (K.cfg.t["alpha_drop"], cond, book.get("version")), ""]
    if not low:
        L.append("No code is below the floor. Nothing to decide.")
    for cid in low + undefined:
        c, r = books[cid], rel.get(cid, {})
        L += ["## %s %s (%s)%s" % (cid, c["label"], c.get("rq", ""), "  [alpha undefined: no variation]" if cid in undefined else ""), "",
              "- alpha %s; assigned by any coder on %s units, at consensus on %s, unanimously on %s" % (
                  (r.get("alpha") or "n/a")[:5], r.get("n_any"), r.get("n_consensus"), r.get("n_unanimous")),
              "- Definition: %s" % c.get("definition", "")]
        if c.get("include"):
            L.append("- Include: %s" % c["include"])
        if c.get("exclude"):
            L.append("- Exclude: %s" % c["exclude"])
        L += ["", "Units the coders disagreed on (reliability units):", ""]
        n = 0
        for u in corpus:
            if u not in rel_units or u not in long:
                continue
            votes = {k: cid in cs for k, (cs, _) in long[u].items()}
            if len(set(votes.values())) > 1:
                n += 1
                L.append("- **%s**: %s" % (u, unit_text(corpus[u]).replace("\n", " ")))
                for k, (cs, why) in sorted(long[u].items()):
                    L.append("  - coder %s (%s): %s :: %s" % (k, K.cfg.coders.get(k, "?"), ", ".join(cs) or "(none)", why))
        if not n:
            L.append("- none")
        L.append("")
    (d / "low_alpha_codes.md").write_text("\n".join(L) + "\n", encoding="utf-8")
    with open(d / "decisions_TEMPLATE.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["code", "label", "alpha", "decision", "new_label", "new_definition", "new_include", "new_exclude", "reason", "decided_by"])
        for cid in low:
            w.writerow([cid, books[cid]["label"], (rel.get(cid, {}).get("alpha") or "")[:5], "", "", "", "", "", "", ""])
    (d / "README.md").write_text(LOW_ALPHA_README.format(n=len(low), rel=K.cfg.d["reliability"]["units"], floor=K.cfg.t["alpha_drop"], cond=cond),
                                 encoding="utf-8")
    print("low-alpha (%s): %d code(s) below %s%s -> %s. STOP: the researchers decide, per code, refine and recode or drop."
          % (cond, len(low), K.cfg.t["alpha_drop"], (" (" + ", ".join(low) + ")") if low else "", d))
    return 0


# --------------------------------------------------------------------------- main
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config")
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("blind-pass")
    p = sub.add_parser("review-dashboard")
    p.add_argument("--cond", choices=["lite", "full"], help="candidate codebook to review (default: the config's mode)")
    p.add_argument("--round", default="r1", help="a label kept with the saved decisions")
    p.add_argument("--proposals", help="optional JSON of proposed edits {code_id: {decision, mergeInto?, rq?, reason, round}}")
    p = sub.add_parser("review-packet")
    p.add_argument("--cond", choices=["lite", "full"])
    p = sub.add_parser("review-to-edits")
    p.add_argument("files", nargs="*", help="the review_<name>.json files downloaded from the dashboard "
                                             "(default: every .json in human/02_author_review/returned/)")
    p.add_argument("--cond", choices=["lite", "full"])
    sub.add_parser("heldout-packet")
    p = sub.add_parser("spot-check")
    p.add_argument("--cond", choices=K.CONDITIONS)
    p = sub.add_parser("low-alpha")
    p.add_argument("--cond", choices=K.CONDITIONS)
    a = ap.parse_args(argv)
    K.init(a.config)
    return {"blind-pass": cmd_blind_pass, "review-dashboard": cmd_review_dashboard, "review-packet": cmd_review_packet,
            "review-to-edits": cmd_review_to_edits, "heldout-packet": cmd_heldout_packet, "spot-check": cmd_spot_check,
            "low-alpha": cmd_low_alpha}[a.cmd](a)


if __name__ == "__main__":
    sys.exit(main() or 0)
